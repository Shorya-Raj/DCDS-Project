"""Make an ordered, deduplicated link manifest from a Word source document."""

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from urllib.parse import urlparse
from xml.etree import ElementTree as ET

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def paragraphs_with_links(path):
    with zipfile.ZipFile(path) as archive:
        document = ET.fromstring(archive.read("word/document.xml"))
        relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
    targets = {
        item.attrib["Id"]: item.attrib.get("Target", "")
        for item in relationships
        if item.attrib.get("Type", "").endswith("/hyperlink")
    }
    for index, paragraph in enumerate(document.iter(W + "p")):
        text_parts = []
        urls = []
        for child in paragraph:
            if child.tag == W + "hyperlink":
                visible = "".join(t.text or "" for t in child.iter(W + "t"))
                target = targets.get(child.attrib.get(R + "id"), "")
                text_parts.append(visible)
                if target:
                    urls.append(target)
            elif child.tag == W + "r":
                text_parts.append("".join(t.text or "" for t in child.iter(W + "t")))
        yield index, "".join(text_parts).strip(), urls


def existing_source_urls(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Tender Log"]
        header_row = None
        source_column = None
        for row in sheet.iter_rows(min_row=1, max_row=20):
            for cell in row:
                if cell.value == "Source Detail URL":
                    header_row, source_column = cell.row, cell.column
                    break
            if source_column:
                break
        if source_column is None:
            raise ValueError("Tender Log has no 'Source Detail URL' header")
        return {
            str(row[0]).strip()
            for row in sheet.iter_rows(
                min_row=header_row + 1,
                min_col=source_column,
                max_col=source_column,
                values_only=True,
            )
            if row[0]
        }
    finally:
        workbook.close()


def build_manifest(docx_path, marker, known_urls):
    marker_pattern = re.compile(r"\b" + re.escape(marker) + r"\b", re.IGNORECASE)
    started = False
    marker_index = None
    category = None
    occurrences = 0
    by_url = {}
    for index, text, urls in paragraphs_with_links(docx_path):
        if not started:
            if not urls and marker_pattern.search(text):
                started = True
                marker_index = index
            continue
        if not urls:
            if text and text.upper() == text and len(text) < 100:
                category = text
            continue
        for url in urls:
            occurrences += 1
            if url not in by_url:
                by_url[url] = {
                    "sequence": len(by_url) + 1,
                    "url": url,
                    "source_document": docx_path.name,
                    "first_paragraph": index,
                    "paragraphs": [],
                    "categories": [],
                    "host": urlparse(url).hostname,
                    "already_recorded": url in known_urls if known_urls is not None else None,
                }
            item = by_url[url]
            item["paragraphs"].append(index)
            if category and category not in item["categories"]:
                item["categories"].append(category)
    if marker_index is None:
        raise ValueError(f"Marker {marker!r} was not found in a standalone paragraph")
    if not by_url:
        raise ValueError("No hyperlinks followed the marker")
    return marker_index, occurrences, list(by_url.values())


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--after-marker", default="NEW")
    parser.add_argument("--existing-workbook", type=Path, help="read only; marks URLs already in Tender Log")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    if args.output.exists() and not args.overwrite:
        parser.error("Output exists; choose another path or pass --overwrite")
    known = existing_source_urls(args.existing_workbook) if args.existing_workbook else None
    marker_index, occurrences, records = build_manifest(args.docx, args.after_marker, known)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="\n") as stream:
        for record in records:
            stream.write(json.dumps(record, ensure_ascii=False) + "\n")
    print(json.dumps({
        "marker_paragraph": marker_index,
        "link_occurrences": occurrences,
        "unique_urls": len(records),
        "repeated_occurrences": occurrences - len(records),
        "already_recorded": sum(r["already_recorded"] is True for r in records),
        "pending": sum(r["already_recorded"] is False for r in records) if known is not None else None,
        "output": str(args.output),
    }, indent=2))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
