"""Validate and normalize third party works tender listing records."""
from __future__ import annotations

import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import urlparse

CATEGORIES = {"Civil works", "Road works", "Infrastructure works"}
MONEY_MULTIPLIERS = {
    "cr": Decimal("10000000"),
    "crore": Decimal("10000000"),
    "lakh": Decimal("100000"),
    "lac": Decimal("100000"),
    "k": Decimal("1000"),
}
SOURCE_KEY_PREFIXES = {
    "MahaTenders official detail page": "mahatenders",
    "Tender247 category listing": "tender247",
}


def make_tender_key(source_type: str, source_record_id: str) -> str:
    """Build the same source-aware key for both databases and the dashboard."""
    prefix = SOURCE_KEY_PREFIXES.get(source_type)
    record_id = str(source_record_id).strip()
    if not prefix:
        raise ValueError(f"Unsupported tender source type: {source_type!r}")
    if not record_id:
        raise ValueError("Each record needs a non-empty source record ID")
    return f"{prefix}:{record_id}"


def _mongo_decimal(value: Decimal | None):
    """Keep INR values exact when writing them to MongoDB."""
    if value is None:
        return None
    from bson.decimal128 import Decimal128

    return Decimal128(value)


def parse_money(value: str | None) -> Decimal | None:
    """Convert amounts such as '3.57 Cr.' and '8.41 Lakh' to INR."""
    if value is None or not str(value).strip() or str(value).strip().casefold() == "refer docs":
        return None
    text = str(value).replace(",", "").strip()
    match = re.fullmatch(r"(\d+(?:\.\d+)?)\s*([A-Za-z]+)\.?", text)
    if not match:
        raise ValueError(f"Unsupported money amount: {value!r}")
    unit = match.group(2).lower()
    if unit not in MONEY_MULTIPLIERS:
        raise ValueError(f"Unsupported money unit: {unit}")
    try:
        return (Decimal(match.group(1)) * MONEY_MULTIPLIERS[unit]).quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise ValueError(f"Invalid money amount: {value!r}") from exc


def normalize_record(raw: dict) -> dict:
    listing_id = str(raw.get("id", "")).strip()
    if not listing_id or not listing_id.isdigit():
        raise ValueError("Each listing needs a numeric Tender247 ID")
    title = str(raw.get("title", "")).strip()
    if not title:
        raise ValueError(f"Listing {listing_id} has no title")
    category = str(raw.get("category", "")).strip()
    if category not in CATEGORIES:
        raise ValueError(f"Listing {listing_id} has unknown category {category!r}")
    location = str(raw.get("location", "")).strip()
    if not location:
        raise ValueError(f"Listing {listing_id} has no location")
    source_url = str(raw.get("sourceUrl", "")).strip()
    if not source_url.startswith("https://www.tender247.com/"):
        raise ValueError(f"Listing {listing_id} has an unexpected source URL")
    try:
        closing_date = date.fromisoformat(str(raw["closingDate"]))
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError(f"Listing {listing_id} has an invalid closing date") from exc
    source_type = "Tender247 category listing"
    return {
        "tender_key": make_tender_key(source_type, listing_id),
        "source_record_id": listing_id,
        "title": title,
        "category": category,
        "location": location,
        "estimated_value_inr": parse_money(raw.get("value")),
        "emd_inr": parse_money(raw.get("emd")),
        "published_date": None,
        "closing_date": closing_date,
        "source_url": source_url,
        "source_type": source_type,
        "raw": raw,
    }


def _cell_text(value) -> str:
    return "" if value is None else str(value).strip()


def _optional_date_text(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _cell_text(value) or None


def _excel_date(value, label: str, tender_id: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Tender {tender_id} has an invalid {label}") from exc


def _excel_amount(value, label: str, tender_id: str) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError(f"Tender {tender_id} has an invalid {label}")
    text = str(value).replace(",", "").replace("₹", "").strip()
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise ValueError(f"Tender {tender_id} has an invalid {label}: {value!r}") from exc


def _official_detail_url(value: str, tender_id: str) -> str:
    url = _cell_text(value)
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname != "mahatenders.gov.in" or parsed.path != "/nicgep/app":
        raise ValueError(f"Tender {tender_id} has an unexpected official detail URL")
    if "FrontEndAdvancedSearchResult" not in parsed.query:
        raise ValueError(f"Tender {tender_id} URL is not an official tender detail page")
    return url


def normalize_mahatenders_row(raw: dict) -> dict:
    """Map one verified MahaTenders Tender Log row to the app's common record shape."""
    tender_id = _cell_text(raw.get("Tender ID"))
    if not tender_id:
        raise ValueError("MahaTenders rows need a Tender ID")
    title = _cell_text(raw.get("Tender Title"))
    organisation = _cell_text(raw.get("Organisation"))
    category = _cell_text(raw.get("Product Category")) or _cell_text(raw.get("Tender Category"))
    location = _cell_text(raw.get("District / Location"))
    if not title or not organisation or not category or not location:
        raise ValueError(f"Tender {tender_id} is missing a required title, organisation, category, or location")
    source_url = _official_detail_url(raw.get("Source Detail URL"), tender_id)
    published_date = _excel_date(raw.get("e-Published Date"), "published date", tender_id)
    closing_date = _excel_date(raw.get("Closing Date"), "closing date", tender_id)
    estimated_value = _excel_amount(raw.get("Estimated Value (INR)"), "estimated value", tender_id)
    emd = _excel_amount(raw.get("EMD (INR)"), "EMD", tender_id)
    source_type = "MahaTenders official detail page"
    return {
        "tender_key": make_tender_key(source_type, tender_id),
        "source_record_id": tender_id,
        "title": title,
        "category": category,
        "location": location,
        "estimated_value_inr": estimated_value,
        "emd_inr": emd,
        "closing_date": closing_date,
        "published_date": published_date,
        "source_url": source_url,
        "source_type": source_type,
        "raw": dict(raw),
    }


def plan_mahatenders_workbook(path: str | Path) -> dict:
    """Map import-ready official rows and report included rows held for missing evidence."""
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Reading MahaTenders workbooks requires openpyxl; install project requirements first") from exc

    workbook = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
    try:
        if "Tender Log" not in workbook.sheetnames:
            raise ValueError("The workbook must contain a 'Tender Log' sheet")
        sheet = workbook["Tender Log"]
        header_values = next(sheet.iter_rows(min_row=8, max_row=8, values_only=True), None)
        if not header_values:
            raise ValueError("The Tender Log header row 8 is missing")
        headers = [_cell_text(value) for value in header_values]
        if "Tender ID" not in headers or "Include?" not in headers:
            raise ValueError("The Tender Log is missing the Tender ID or Include? column")

        records: list[dict] = []
        held: list[dict] = []
        excluded_count = 0
        seen_ids: set[str] = set()
        for excel_row, values in enumerate(sheet.iter_rows(min_row=9, values_only=True), start=9):
            if not values or values[0] is None:
                continue
            raw = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
            tender_id = _cell_text(raw.get("Tender ID"))
            if not tender_id:
                if _cell_text(raw.get("Include?")) == "Yes":
                    raise ValueError(f"Included row {excel_row} has no Tender ID")
                excluded_count += 1
                continue
            if tender_id in seen_ids:
                raise ValueError(f"Duplicate Tender ID in Tender Log: {tender_id}")
            seen_ids.add(tender_id)
            if _cell_text(raw.get("Include?")) != "Yes":
                excluded_count += 1
                continue

            reasons = []
            if _cell_text(raw.get("Collection Status")) != "Details captured":
                reasons.append("detail page not captured")
            if not _cell_text(raw.get("Source Detail URL")):
                reasons.append("exact source detail URL missing")
            for field in ("Tender Title", "Organisation", "District / Location", "e-Published Date", "Closing Date"):
                if raw.get(field) is None or _cell_text(raw.get(field)) == "":
                    reasons.append(f"{field} missing")
            if reasons:
                held.append({"excel_row": excel_row, "tender_id": tender_id, "reasons": reasons})
                continue

            records.append(normalize_mahatenders_row(raw))
        if len({record["tender_key"] for record in records}) != len(records):
            raise ValueError("The import plan contains duplicate Tender IDs")
        return {"records": records, "held": held, "excluded_count": excluded_count}
    finally:
        workbook.close()


def load_records(path: str | Path) -> list[dict]:
    source_path = Path(path)
    if source_path.suffix.casefold() in {".xlsx", ".xlsm"}:
        return plan_mahatenders_workbook(source_path)["records"]
    with Path(path).open(encoding="utf-8-sig") as source:
        payload = json.load(source)
    if not isinstance(payload, list):
        raise ValueError("The source file must contain a JSON array")
    records = [normalize_record(item) for item in payload]
    keys = [record["tender_key"] for record in records]
    if len(keys) != len(set(keys)):
        raise ValueError("The source file contains duplicate source record IDs")
    return records


def document_for(record: dict) -> dict:
    """Keep variable listing details in MongoDB without inventing missing fields."""
    raw = record["raw"]
    common = {
        "_id": record["tender_key"],
        "source_record_id": record["source_record_id"],
        "source_type": record["source_type"],
        "title": record["title"],
        "category": record["category"],
        "categories": raw.get("categories") or [record["category"]],
        "location": record["location"],
        "published_date": record["published_date"].isoformat() if record["published_date"] else None,
        "closing_date": record["closing_date"].isoformat(),
        "estimated_value_inr": _mongo_decimal(record["estimated_value_inr"]),
        "emd_inr": _mongo_decimal(record["emd_inr"]),
        "source_url": record["source_url"],
    }
    if record.get("source_type") == "MahaTenders official detail page":
        return {
            **common,
            "reference_no": raw.get("Tender Reference No."),
            "organisation": raw.get("Organisation"),
            "department": raw.get("Department / Division"),
            "tender_category": raw.get("Tender Category"),
            "product_category": raw.get("Product Category"),
            "opening_date": _optional_date_text(raw.get("Opening Date")),
            "portal_status_snapshot": raw.get("Portal Status"),
            "status_collected_on": _optional_date_text(raw.get("Collected On")),
            "tender_fee_inr": raw.get("Tender Fee (INR)"),
            "work_period_days": raw.get("Work Period (Days)"),
            "scope": raw.get("Scope / Work Description"),
            "eligibility": raw.get("Eligibility / Conditions"),
            "corrigendum_snapshot": raw.get("Corrigendum?"),
            "document_names_or_links": raw.get("Document Names / Links"),
            "notes": raw.get("Notes / Missing Fields"),
            "raw_source_row": raw,
        }
    return {
        **common,
        "raw_value": raw.get("value"),
        "raw_emd": raw.get("emd"),
        "raw_countdown": raw.get("countdown"),
        "eligibility": None,
        "document_references": None,
        "corrigenda": None,
    }

